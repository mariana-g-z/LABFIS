"""
data_manager.py
---------------
Módulo central de acceso a los archivos xlsx/csv.
Todos los backends importan de aquí para leer y escribir datos persistentes.
Los datos se cargan al arrancar y se guardan en cada operación de escritura.
"""
import os
import csv
import hashlib
from datetime import datetime
from typing import Dict, List, Optional, Any

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# ── Rutas por defecto (relativas a la raíz del proyecto) ──────────────────────
DEFAULT_PERSONAS_XLSX   = "personas.xlsx"
DEFAULT_MATERIALES_XLSX = "materiales.xlsx"
DEFAULT_PRESTAMOS_XLSX  = "prestamos.xlsx"
DEFAULT_BAJAS_XLSX      = "bajas.xlsx"

# ── Singleton de rutas activas (se actualizan desde Configuración) ─────────────
_paths: Dict[str, str] = {
    "personas":   DEFAULT_PERSONAS_XLSX,
    "materiales": DEFAULT_MATERIALES_XLSX,
    "prestamos":  DEFAULT_PRESTAMOS_XLSX,
    "bajas":      DEFAULT_BAJAS_XLSX,
}

def set_paths(personas: str, materiales: str, prestamos: str, bajas: str):
    """Actualiza las rutas activas. Llamado desde ConfiguracionBackend al importar."""
    _paths["personas"]   = personas
    _paths["materiales"] = materiales
    _paths["prestamos"]  = prestamos
    _paths["bajas"]      = bajas
    # Recargar datos en todos los backends
    _reload_all()

def get_paths() -> Dict[str, str]:
    return dict(_paths)

# ── Helpers de lectura xlsx ────────────────────────────────────────────────────
def _read_xlsx(path: str) -> List[Dict[str, Any]]:
    """Lee un xlsx y devuelve lista de dicts (primera fila = cabeceras)."""
    if not OPENPYXL_AVAILABLE or not os.path.exists(path):
        return []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [str(h) if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        result = []
        for row in rows[1:]:
            if all(v is None for v in row):
                continue
            result.append({headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))})
        return result
    except Exception as e:
        print(f"[DataManager] Error leyendo {path}: {e}")
        return []

def _write_xlsx(path: str, headers: List[str], rows: List[List[Any]]):
    """Escribe (sobreescribe) un xlsx con headers y filas dadas."""
    if not OPENPYXL_AVAILABLE:
        # Fallback: escribir CSV
        csv_path = path.replace(".xlsx", ".csv")
        try:
            with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(headers)
                w.writerows(rows)
        except Exception as e:
            print(f"[DataManager] Error escribiendo CSV {csv_path}: {e}")
        return
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append([str(v) if v is not None else "" for v in row])
        wb.save(path)
        # También exportar CSV espejo
        csv_path = path.replace(".xlsx", ".csv")
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(headers)
            w.writerows([[str(v) if v is not None else "" for v in row] for row in rows])
    except Exception as e:
        print(f"[DataManager] Error escribiendo {path}: {e}")

# ── Caché en memoria ──────────────────────────────────────────────────────────
_personas_cache:   List[Dict] = []
_materiales_cache: List[Dict] = []
_prestamos_cache:  List[Dict] = []
_bajas_cache:      List[Dict] = []
_loaded = False

def _load_all():
    global _personas_cache, _materiales_cache, _prestamos_cache, _bajas_cache, _loaded
    _personas_cache   = _read_xlsx(_paths["personas"])
    _materiales_cache = _read_xlsx(_paths["materiales"])
    _prestamos_cache  = _read_xlsx(_paths["prestamos"])
    _bajas_cache      = _read_xlsx(_paths["bajas"])
    _loaded = True

# Callbacks registrados por backends para recargar cuando cambien rutas
_reload_callbacks = []

def register_reload_callback(fn):
    """Los backends llaman esto en su __init__ para ser notificados al recargar."""
    if fn not in _reload_callbacks:
        _reload_callbacks.append(fn)

def _reload_all():
    global _loaded
    _loaded = False
    _load_all()
    # Notificar a todos los backends registrados
    for fn in _reload_callbacks:
        try:
            fn()
        except Exception as e:
            print(f"[DataManager] Error en callback de recarga: {e}")

def _ensure_loaded():
    if not _loaded:
        _load_all()

# ── API de Personas ───────────────────────────────────────────────────────────
PERSONAS_HEADERS = ["username","full_name","role","career","email","notes","status","password_hash","active_loans","created_at"]

def get_all_personas() -> List[Dict]:
    _ensure_loaded()
    return list(_personas_cache)

def get_persona(username: str) -> Optional[Dict]:
    _ensure_loaded()
    for p in _personas_cache:
        if str(p.get("username","")) == username:
            return dict(p)
    return None

def get_estudiantes() -> List[Dict]:
    _ensure_loaded()
    return [p for p in _personas_cache if str(p.get("role","")).lower() == "estudiante"]

def get_usuarios_sistema() -> List[Dict]:
    _ensure_loaded()
    return [p for p in _personas_cache if str(p.get("role","")).lower() in ("administrador","becario")]

def save_persona(data: Dict):
    """Inserta o actualiza una persona por username."""
    _ensure_loaded()
    username = str(data.get("username",""))
    for i, p in enumerate(_personas_cache):
        if str(p.get("username","")) == username:
            _personas_cache[i] = dict(data)
            _flush_personas()
            return
    _personas_cache.append(dict(data))
    _flush_personas()

def delete_persona(username: str):
    global _personas_cache
    _ensure_loaded()
    _personas_cache = [p for p in _personas_cache if str(p.get("username","")) != username]
    _flush_personas()

def _flush_personas():
    rows = []
    for p in _personas_cache:
        rows.append([p.get(h,"") for h in PERSONAS_HEADERS])
    _write_xlsx(_paths["personas"], PERSONAS_HEADERS, rows)

# ── API de Materiales ─────────────────────────────────────────────────────────
MATERIALES_HEADERS = ["code","name","category","control_type","stock","status","description","created_at","folio_baja"]

def get_all_materiales() -> List[Dict]:
    _ensure_loaded()
    return list(_materiales_cache)

def get_material(code: str) -> Optional[Dict]:
    _ensure_loaded()
    for m in _materiales_cache:
        if str(m.get("code","")).upper() == code.upper():
            return dict(m)
    return None

def save_material(data: Dict):
    _ensure_loaded()
    code = str(data.get("code","")).upper()
    for i, m in enumerate(_materiales_cache):
        if str(m.get("code","")).upper() == code:
            _materiales_cache[i] = dict(data)
            _flush_materiales()
            return
    _materiales_cache.append(dict(data))
    _flush_materiales()

def delete_material(code: str):
    global _materiales_cache
    _ensure_loaded()
    _materiales_cache = [m for m in _materiales_cache if str(m.get("code","")).upper() != code.upper()]
    _flush_materiales()

def _flush_materiales():
    rows = []
    for m in _materiales_cache:
        rows.append([m.get(h,"") for h in MATERIALES_HEADERS])
    _write_xlsx(_paths["materiales"], MATERIALES_HEADERS, rows)

# ── API de Préstamos ──────────────────────────────────────────────────────────
PRESTAMOS_HEADERS = ["loan_id","student_id","student_name","material_code","material_name","quantity","loan_date","due_date","return_date","status","becario"]

def get_all_prestamos() -> List[Dict]:
    _ensure_loaded()
    return list(_prestamos_cache)

def get_prestamo(loan_id: str) -> Optional[Dict]:
    _ensure_loaded()
    for p in _prestamos_cache:
        if str(p.get("loan_id","")) == loan_id:
            return dict(p)
    return None

def save_prestamo(data: Dict):
    _ensure_loaded()
    loan_id = str(data.get("loan_id",""))
    for i, p in enumerate(_prestamos_cache):
        if str(p.get("loan_id","")) == loan_id:
            _prestamos_cache[i] = dict(data)
            _flush_prestamos()
            return
    _prestamos_cache.append(dict(data))
    _flush_prestamos()

def _flush_prestamos():
    rows = []
    for p in _prestamos_cache:
        rows.append([p.get(h,"") for h in PRESTAMOS_HEADERS])
    _write_xlsx(_paths["prestamos"], PRESTAMOS_HEADERS, rows)

def get_next_loan_id() -> str:
    _ensure_loaded()
    ids = []
    for p in _prestamos_cache:
        lid = str(p.get("loan_id",""))
        if lid.startswith("P") and lid[1:].isdigit():
            ids.append(int(lid[1:]))
    return f"P{(max(ids)+1) if ids else 100}"

# ── API de Bajas ──────────────────────────────────────────────────────────────
BAJAS_HEADERS = ["fecha","codigo","nombre","tipo","folio","usuario","motivo"]

def get_all_bajas() -> List[Dict]:
    _ensure_loaded()
    return list(_bajas_cache)

def save_baja(data: Dict):
    _ensure_loaded()
    _bajas_cache.append(dict(data))
    rows = []
    for b in _bajas_cache:
        rows.append([b.get(h,"") for h in BAJAS_HEADERS])
    _write_xlsx(_paths["bajas"], BAJAS_HEADERS, rows)

# ── Hash de contraseña ────────────────────────────────────────────────────────
def hash_password(password: str) -> str:
    return hashlib.sha256(password.encode()).hexdigest()

def verify_password(password: str, stored_hash: str) -> bool:
    return hash_password(password) == str(stored_hash)
